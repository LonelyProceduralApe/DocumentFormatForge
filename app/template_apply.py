"""版头模板套用：以源文档为底，把模板版头（红头/发文字号/红线）插入到开头。

版头为 VML 矢量图形/文本（无媒体文件、自包含），深拷贝后插入源文档开头即可，
源文档的横向节、图片、表格、分页符等原样保留。

支持：
- 模板内置打包进 exe（bundled_template_dir），同时允许用户在 exe 旁再放 template 目录增补/覆盖；
- 填写模板里的占位内容（发文字号、签发人姓名）。
"""
from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import formatter
from .docx_io import set_indent_chars, set_line_exact, set_run_font

TEMPLATE_DIR_NAME = "template"


def _meipass() -> Path:
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent.parent))


def bundled_template_dir() -> Path:
    """内置模板目录（打包时随 exe 附带）。"""
    return _meipass() / TEMPLATE_DIR_NAME


def user_template_dir() -> Path:
    """用户模板目录（exe 同目录的 template\，可增补/覆盖内置模板）。"""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent / TEMPLATE_DIR_NAME
    return Path(__file__).resolve().parent.parent / TEMPLATE_DIR_NAME


def list_templates() -> list:
    """返回 [(显示名, 路径)]；内置 + 用户模板合并，同名以用户目录为准。"""
    result = {}
    for d in (bundled_template_dir(), user_template_dir()):
        if d.is_dir():
            for f in sorted(d.glob("*.docx")):
                if f.name.startswith("~$"):
                    continue  # 跳过 Word/WPS 锁文件
                result[f.stem] = f
    return [(name, path) for name, path in result.items()]


def extract_issuer_fields(template_path) -> dict:
    """提取模板里「发文字号」与「签发人姓名」占位文本，供界面预填。"""
    try:
        doc = Document(str(template_path))
        for p in doc.paragraphs:
            t = p.text or ""
            if "签发人" in t:
                left, right = t.split("签发人", 1)
                return {"doc_num": left.strip(), "name": right.lstrip("：").strip()}
    except Exception:
        pass
    return {"doc_num": "", "name": ""}


def fill_issuer_line(document, doc_num: str, name: str, ascii_font: str = "Times New Roman") -> bool:
    """填写发文字号/签发人。

    - 都未填：发文字号行只保留红线 pict（清除发文字号、签发人等文本）；
    - 只填发文字号：发文字号居中，删除签发人；
    - 都填：发文字号左空一字、签发人右空一字（右制表位）。
    """
    doc_num = (doc_num or "").strip()
    name = (name or "").strip()
    for p in document.paragraphs:
        if "签发人" in (p.text or ""):
            _rebuild_issuer_para(p, doc_num, name, ascii_font, document)
            return True
    return False


def _is_pict_run(run) -> bool:
    return bool(run._r.findall(".//" + qn("w:drawing"))) or \
        bool(run._r.findall(".//" + qn("w:pict")))


def _rebuild_issuer_para(p, doc_num: str, name: str, ascii_font: str, document) -> None:
    size = 16.0
    runs = list(p.runs)
    pict_els = [r._r for r in runs if _is_pict_run(r)]
    for r in runs:
        r._r.getparent().remove(r._r)

    if not doc_num and not name:
        # 都未填：只保留红线 pict，清除全部文本
        for el in pict_els:
            p._p.append(el)
        return

    if name:
        # 发文字号左空一字，签发人右空一字（右制表位）
        set_indent_chars(p, size, left_chars=1, right_chars=1)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _add_right_tab(p, document, size)
    else:
        # 只填发文字号：居中，删除签发人
        set_indent_chars(p, size, left_chars=0, right_chars=0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 先放红线 pict（保持模板顺序），再放文本
    for el in pict_els:
        p._p.append(el)

    _add_run(p, doc_num, "仿宋_GB2312", ascii_font, size)
    if name:
        _add_tab_run(p, ascii_font, size)
        _add_run(p, "签发人：", "仿宋_GB2312", ascii_font, size)
        _add_run(p, name, "楷体_GB2312", ascii_font, size)


def _add_run(p, text, east, ascii_font, size):
    run = p.add_run(text)
    set_run_font(run, east, ascii_font, size)
    return run


def _add_tab_run(p, ascii_font, size):
    run = p.add_run()
    run._r.append(OxmlElement("w:tab"))
    set_run_font(run, "仿宋_GB2312", ascii_font, size)
    return run


def _add_right_tab(p, document, size):
    sec = document.sections[0]
    text_width_twips = sec.page_width.twips - sec.left_margin.twips - sec.right_margin.twips
    pos = text_width_twips - int(size * 20)  # 右空一字
    pPr = p._p.get_or_add_pPr()
    tabs = pPr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        pPr.append(tabs)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(pos))
    tabs.append(tab)


def _has_content(p) -> bool:
    if (p.text or "").strip():
        return True
    if p._p.findall(".//" + qn("w:drawing")) or p._p.findall(".//" + qn("w:pict")):
        return True
    return False


def apply_template(template_path, src_doc, roles: dict, rules: dict,
                   doc_num=None, issuer_name=None, date_text=None):
    """以源文档为底，套用版头与版记，返回 Document（未保存）。

    doc_num / issuer_name / date_text：发文字号、签发人姓名、日期（成文日期与印发日期）。
    """
    tpl = Document(str(template_path))

    # 1) 提取模板版头（红头～红线）与版记（黑线～印发）
    header_els = _extract_header_elements(tpl)
    imprint_el = _extract_imprint_element(tpl)
    imprint_fields = extract_imprint_fields(template_path)

    # 2) 补齐 VML 等命名空间
    _ensure_namespaces(src_doc, tpl)

    # 3) 对源文档套格式（含删除版头/版记、替换成文日期）
    formatter.apply_formatting(src_doc, roles, rules, date_text)

    # 4) 插入版头
    body_line = rules.get("line_spacing_pt", {}).get("body", 29)
    _insert_header(src_doc, header_els, body_line)

    # 5) 填写发文字号/签发人
    ascii_font = rules.get("fonts", {}).get("ascii", "Times New Roman")
    fill_issuer_line(src_doc, doc_num or "", issuer_name or "", ascii_font)

    # 6) 追加版记
    if imprint_el is not None:
        print_org = imprint_fields.get("org", "")
        placeholder = imprint_fields.get("print_date", "")
        if date_text:
            print_date = date_text + "印发"
        else:
            print_date = placeholder
        _apply_imprint(src_doc, imprint_el, print_org, print_date, ascii_font)

    return src_doc


def _extract_header_elements(tpl):
    """提取版头：第 0 段到发文字号行（含"签发人"或"〔年〕号"）为止。"""
    paras = list(tpl.paragraphs)
    last = -1
    for i, p in enumerate(paras):
        last = i
        t = p.text or ""
        if "签发人" in t or ("〔" in t and "号" in t):
            break
    if last == -1:
        return []
    return [copy.deepcopy(paras[i]._p) for i in range(0, last + 1)]


def _extract_imprint_element(tpl):
    """深拷贝模板版记段落（最后一个有内容的段落），保留黑线、行距、字体等。"""
    paras = list(tpl.paragraphs)
    last = -1
    for i, p in enumerate(paras):
        if _has_content(p):
            last = i
    if last == -1:
        return None
    return copy.deepcopy(paras[last]._p)


def extract_imprint_fields(template_path) -> dict:
    """提取版记里的印发机关与印发日期占位文本。"""
    try:
        doc = Document(str(template_path))
        for p in doc.paragraphs:
            t = (p.text or "").strip()
            if not t:
                continue
            if "印发" in t or "办公室" in t or "委员会" in t or "政府" in t:
                m = re.search(
                    r"([xX×0-9]{1,4}\s*年\s*[xX×0-9]{1,2}\s*月\s*[xX×0-9]{1,2}\s*日(?:\s*印发)?)",
                    t,
                )
                if m:
                    return {"org": t[:m.start()].strip(), "print_date": m.group(1).strip()}
                return {"org": t, "print_date": ""}
    except Exception:
        pass
    return {"org": "", "print_date": ""}


def _apply_imprint(src_doc, imprint_el, print_org, print_date, ascii_font):
    """把版记（深拷贝原段落）追加到末尾，重建文字为 印发机关+制表符+印发日期。"""
    body = src_doc.element.body
    sectPr = body.find(qn("w:sectPr"))
    if sectPr is not None:
        sectPr.addprevious(imprint_el)
    else:
        body.append(imprint_el)
    for p in src_doc.paragraphs:
        if p._p is imprint_el:
            _rebuild_imprint_para(p, print_org, print_date, ascii_font, src_doc)
            break


def _rebuild_imprint_para(p, print_org, print_date, ascii_font, document):
    size = 14.0
    runs = list(p.runs)
    pict_els = [r._r for r in runs if _is_pict_run(r)]
    for r in runs:
        r._r.getparent().remove(r._r)

    # 段落属性：两端对齐、左右各缩进1字符、特殊格式无、段前段后0、29磅行距
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_indent_chars(p, size, left_chars=1, right_chars=1)
    set_line_exact(p, 29)
    _add_right_tab(p, document, size)

    for el in pict_els:
        p._p.append(el)

    _add_run(p, print_org or "", "仿宋_GB2312", ascii_font, size)
    if print_date:
        _add_tab_run(p, ascii_font, size)
        _add_run(p, print_date, "仿宋_GB2312", ascii_font, size)


def _add_page_break(document):
    p = document.add_paragraph()
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)
    return p


def _ensure_namespaces(src_doc, tpl):
    """把模板根元素上的命名空间补齐到源文档根元素，确保 VML 正常渲染。"""
    src_root = src_doc.element
    for prefix, uri in (tpl.element.nsmap or {}).items():
        if prefix and uri and prefix not in (src_root.nsmap or {}):
            try:
                src_root.set("xmlns:" + prefix, uri)
            except Exception:
                pass


def _insert_header(src_doc, header_els, line_pt: float = 29):
    """在源文档第一个段落前插入版头，并在其后加两个空段（红线下方空二行）。"""
    body = src_doc.element.body
    first_p = body.find(qn("w:p"))
    if first_p is None:
        for el in header_els:
            body.append(el)
        body.append(_make_empty_p(line_pt))
        body.append(_make_empty_p(line_pt))
        return
    for el in header_els:
        first_p.addprevious(el)
    for _ in range(2):
        first_p.addprevious(_make_empty_p(line_pt))


def _make_empty_p(line_pt: float):
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:line"), str(int(round(line_pt * 20))))
    spacing.set(qn("w:lineRule"), "exact")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    pPr.append(spacing)
    p.append(pPr)
    return p
