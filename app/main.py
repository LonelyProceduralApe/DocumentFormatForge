"""公文格式自动排版助手 - 主窗口与入口。"""
from __future__ import annotations

import sys
from pathlib import Path

# 开发环境下把工作区 .pylib 加入导入路径（python-docx 装于此处时可用）
_ROOT = Path(__file__).resolve().parent.parent
_PYLIB = _ROOT / ".pylib"
if _PYLIB.is_dir() and str(_PYLIB) not in sys.path:
    sys.path.insert(0, str(_PYLIB))

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QApplication,
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QDoubleSpinBox,
    QFileDialog,
    QFormLayout,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QInputDialog,
    QLabel,
    QLineEdit,
    QListWidget,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from docx import Document

from . import classifier, font_check, formatter, rules_config, template_apply, wps_launch
from .docx_io import read_paragraphs


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.VERSION = "v2.1.2"
        self.setWindowTitle(f"公文格式自动排版助手 {self.VERSION}    by  LonelyProceduralApe")
        self.resize(1120, 760)

        self.rules = rules_config.load_rules()
        self.files: list[Path] = []
        self.current_path: Path | None = None
        self.current_paras: list[dict] = []
        self.current_roles: dict[int, str] = {}

        self._build_ui()
        self._check_fonts_on_start()

    # ---------- UI ----------
    def _build_ui(self):
        splitter = QSplitter(Qt.Horizontal)

        # 左侧：文件列表
        left = QWidget()
        left_layout = QVBoxLayout(left)
        group = QGroupBox("待排版文件")
        gl = QVBoxLayout(group)
        self.file_list = QListWidget()
        self.file_list.currentRowChanged.connect(lambda _r: self.load_selected())
        gl.addWidget(self.file_list)
        row1 = QHBoxLayout()
        self.btn_add = QPushButton("添加文件…")
        self.btn_add.clicked.connect(self.add_files)
        self.btn_remove = QPushButton("移除选中")
        self.btn_remove.clicked.connect(self.remove_selected)
        self.btn_clear = QPushButton("清空")
        self.btn_clear.clicked.connect(self.clear_files)
        row1.addWidget(self.btn_add)
        row1.addWidget(self.btn_remove)
        row1.addWidget(self.btn_clear)
        gl.addLayout(row1)
        left_layout.addWidget(group)
        tpl_row = QHBoxLayout()
        tpl_row.addWidget(QLabel("版头模板:"))
        self.template_combo = QComboBox()
        tpl_row.addWidget(self.template_combo, 1)
        left_layout.addLayout(tpl_row)
        self.docnum_edit = QLineEdit()
        self.docnum_edit.setPlaceholderText("如：白委〔2026〕12号")
        self.issuer_edit = QLineEdit()
        self.issuer_edit.setPlaceholderText("如：张三（无签发人可留空）")
        self.date_edit = QLineEdit()
        self.date_edit.setPlaceholderText("如：2024年1月5日（成文日期/印发日期）")
        issuer_form = QFormLayout()
        issuer_form.addRow("发文字号:", self.docnum_edit)
        issuer_form.addRow("签发人:", self.issuer_edit)
        issuer_form.addRow("日期:", self.date_edit)
        left_layout.addLayout(issuer_form)
        self.template_combo.currentIndexChanged.connect(lambda _i: self._on_template_changed())
        self._refresh_templates()
        left_layout.addWidget(self._btn("套用并另存", self.apply_selected))
        left_layout.addWidget(self._btn("全部套用（自动）", self.apply_all))
        left_layout.addStretch(1)

        # 右侧：段落清单
        right = QWidget()
        right_layout = QVBoxLayout(right)
        self.cur_label = QLabel("请先添加并选中一个 .docx 文件")
        right_layout.addWidget(self.cur_label)
        self.table = QTableWidget(0, 4)
        self.table.setHorizontalHeaderLabels(["序号", "段落内容", "角色", "将套用格式"])
        self.table.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeToContents)
        self.table.verticalHeader().setVisible(False)
        right_layout.addWidget(self.table)
        row2 = QHBoxLayout()
        row2.addWidget(self._btn("重新自动识别", self.reclassify))
        row2.addWidget(self._btn("规则管理", self.open_rule_manager))
        row2.addWidget(self._btn("规则设置", self.open_rules_dialog))
        row2.addWidget(self._btn("字体检查", self.check_fonts_dialog))
        row2.addStretch(1)
        right_layout.addLayout(row2)

        splitter.addWidget(left)
        splitter.addWidget(right)
        splitter.setSizes([280, 840])
        self.setCentralWidget(splitter)

        self.statusBar().showMessage("就绪")

    def _btn(self, text, handler):
        b = QPushButton(text)
        b.clicked.connect(handler)
        return b

    def _refresh_templates(self):
        self.template_combo.clear()
        self.template_combo.addItem("（不添加版头）", None)
        for name, path in template_apply.list_templates():
            self.template_combo.addItem(name, str(path))

    def _selected_template_path(self):
        data = self.template_combo.currentData() if self.template_combo else None
        return Path(data) if data else None

    def _on_template_changed(self):
        tpl = self._selected_template_path()
        if tpl is None:
            self.docnum_edit.setText("")
            self.issuer_edit.setText("")
            self.docnum_edit.setEnabled(False)
            self.issuer_edit.setEnabled(False)
            return
        self.docnum_edit.setEnabled(True)
        self.issuer_edit.setEnabled(True)
        try:
            fields = template_apply.extract_issuer_fields(tpl)
            self.docnum_edit.setText(fields.get("doc_num", ""))
            self.issuer_edit.setText(fields.get("name", ""))
        except Exception:
            pass

    def _issuer_values(self):
        tpl = self._selected_template_path()
        if tpl is None:
            return None, None
        return self.docnum_edit.text().strip(), self.issuer_edit.text().strip()

    # ---------- 文件列表 ----------
    def add_files(self):
        paths, _ = QFileDialog.getOpenFileNames(self, "选择 docx 文件", "", "Word 文档 (*.docx)")
        for p in paths:
            if p:
                fp = Path(p)
                if fp not in self.files:
                    self.files.append(fp)
        self._refresh_file_list()

    def remove_selected(self):
        row = self.file_list.currentRow()
        if 0 <= row < len(self.files):
            self.files.pop(row)
            self._refresh_file_list()

    def clear_files(self):
        self.files.clear()
        self.current_path = None
        self.current_paras = []
        self.current_roles = {}
        self.table.setRowCount(0)
        self._refresh_file_list()

    def _refresh_file_list(self):
        self.file_list.clear()
        for f in self.files:
            self.file_list.addItem(str(f))

    # ---------- 加载与识别 ----------
    def load_selected(self):
        item = self.file_list.currentItem()
        if item is None:
            return
        path = Path(item.text())
        try:
            doc = Document(str(path))
        except Exception as exc:  # noqa: BLE001
            QMessageBox.warning(self, "打开失败", f"{path.name}\n{exc}")
            return
        self.current_path = path
        self.current_paras = read_paragraphs(doc)
        self.current_roles = classifier.classify(self.current_paras)
        self.cur_label.setText(f"当前：{path.name}（共 {len(self.current_paras)} 段）")
        self.populate_table()

    def reclassify(self):
        if self.current_path is None:
            return
        self.current_roles = classifier.classify(self.current_paras)
        self.populate_table()

    def populate_table(self):
        self.table.setRowCount(0)
        for i, p in enumerate(self.current_paras):
            row = self.table.rowCount()
            self.table.insertRow(row)
            idx_item = QTableWidgetItem(str(p["idx"] + 1))
            idx_item.setFlags(Qt.ItemIsEnabled)
            text_item = QTableWidgetItem(p["text"] if p["text"] else "（空行）")
            text_item.setFlags(Qt.ItemIsEnabled)
            self.table.setItem(row, 0, idx_item)
            self.table.setItem(row, 1, text_item)

            combo = QComboBox()
            combo.addItem("（跳过/不处理）", None)
            combo.addItem("（删除）", "delete")
            for role in rules_config.ROLES:
                combo.addItem(rules_config.ROLE_LABELS[role], role)
            role = self.current_roles.get(p["idx"])
            if role is None:
                combo.setCurrentIndex(0)
            else:
                data_idx = combo.findData(role)
                combo.setCurrentIndex(data_idx if data_idx >= 0 else 0)
            combo.currentIndexChanged.connect(lambda _i=0, row=row: self._on_role_changed(row))
            self.table.setCellWidget(row, 2, combo)

            self.table.setItem(row, 3, QTableWidgetItem(self._format_desc(role)))

    def _on_role_changed(self, row):
        combo = self.table.cellWidget(row, 2)
        role = combo.currentData() if combo else None
        if row < len(self.current_paras):
            para_idx = self.current_paras[row]["idx"]
            if role is None:
                self.current_roles.pop(para_idx, None)
            else:
                self.current_roles[para_idx] = role
            item = QTableWidgetItem(self._format_desc(role))
            item.setFlags(Qt.ItemIsEnabled)
            self.table.setItem(row, 3, item)

    def _format_desc(self, role):
        if role is None:
            return "保持原样"
        if role == "delete":
            return "删除该段落"
        r = self.rules
        return (f"{rules_config.role_font(r, role)} "
                f"{rules_config.role_size_pt(r, role):g}pt "
                f"行距{rules_config.role_line_spacing_pt(r, role):g}pt "
                f"{rules_config.role_alignment(r, role)}")

    def _collect_roles_from_table(self):
        roles = {}
        for row in range(self.table.rowCount()):
            combo = self.table.cellWidget(row, 2)
            if combo is None or row >= len(self.current_paras):
                continue
            role = combo.currentData()
            if role is not None:
                roles[self.current_paras[row]["idx"]] = role
        self.current_roles = roles

    # ---------- 处理 ----------
    def _process(self, src: Path, roles: dict) -> Path:
        doc = Document(str(src))
        date_text = self.date_edit.text().strip() or None
        tpl = self._selected_template_path()
        if tpl is not None:
            doc_num, name = self._issuer_values()
            out_doc = template_apply.apply_template(tpl, doc, roles, self.rules,
                                                    doc_num=doc_num, issuer_name=name,
                                                    date_text=date_text)
        else:
            formatter.apply_formatting(doc, roles, self.rules, date_text)
            out_doc = doc
        out = self._output_path(src)
        out_doc.save(str(out))
        return out

    def _output_path(self, src: Path) -> Path:
        suffix = self.rules.get("output", {}).get("suffix", "_已排版")
        return src.with_name(f"{src.stem}{suffix}.docx")

    def apply_selected(self):
        if self.current_path is None:
            QMessageBox.information(self, "提示", "请先选择并加载一个文件")
            return
        self._collect_roles_from_table()
        try:
            out = self._process(self.current_path, self.current_roles)
        except Exception as exc:  # noqa: BLE001
            QMessageBox.critical(self, "处理失败", str(exc))
            return
        QMessageBox.information(self, "完成", f"已生成：\n{out}")
        wps_launch.open_with_wps(out)

    def apply_all(self):
        if not self.files:
            QMessageBox.information(self, "提示", "请先添加文件")
            return
        outs = []
        errors = []
        for f in self.files:
            try:
                doc = Document(str(f))
                roles = classifier.classify(read_paragraphs(doc))
                date_text = self.date_edit.text().strip() or None
                tpl = self._selected_template_path()
                if tpl is not None:
                    doc_num, name = self._issuer_values()
                    out_doc = template_apply.apply_template(tpl, doc, roles, self.rules,
                                                            doc_num=doc_num, issuer_name=name,
                                                            date_text=date_text)
                else:
                    formatter.apply_formatting(doc, roles, self.rules, date_text)
                    out_doc = doc
                out = self._output_path(f)
                out_doc.save(str(out))
                outs.append(out)
            except Exception as exc:  # noqa: BLE001
                errors.append(f"{f.name}: {exc}")
        if errors:
            QMessageBox.warning(self, "部分失败", "\n".join(errors))
        if outs:
            QMessageBox.information(self, "完成", f"已生成 {len(outs)} 个文件")
            wps_launch.open_with_wps(outs[0])
            try:
                import os
                os.startfile(str(outs[0].parent))
            except Exception:
                pass

    # ---------- 规则与字体 ----------
    def open_rule_manager(self):
        dlg = RuleManagerDialog(self.rules, self)
        if dlg.exec() == QDialog.Accepted:
            self.rules = dlg.get_rules()
            rules_config.save_rules(self.rules)
            self.populate_table()

    def open_rules_dialog(self):
        dlg = RulesDialog(self.rules, self)
        if dlg.exec() == QDialog.Accepted:
            self.rules = dlg.get_rules()
            rules_config.save_rules(self.rules)
            self.populate_table()

    def _check_fonts_on_start(self):
        missing = font_check.check_fonts(set(self.rules["fonts"].values()))
        if missing:
            self.statusBar().showMessage("⚠ 缺少字体：" + "、".join(missing))
        else:
            self.statusBar().showMessage("字体齐全")

    def check_fonts_dialog(self):
        missing = font_check.check_fonts(set(self.rules["fonts"].values()))
        if missing:
            QMessageBox.warning(self, "字体检查", "以下字体未检测到（可能回退显示）：\n" + "\n".join(missing))
        else:
            QMessageBox.information(self, "字体检查", "规则所需字体均已安装。")


class RulesDialog(QDialog):
    """规则快速编辑对话框；更完整的修改可点“打开 rules.json”。"""

    def __init__(self, rules, parent=None):
        super().__init__(parent)
        self.setWindowTitle("规则设置")
        self.setMinimumWidth(420)
        self.rules = rules
        form = QFormLayout(self)
        f = rules["fonts"]
        s = rules["sizes_pt"]
        l = rules["line_spacing_pt"]
        pg = rules["page"]

        self.e_title_font = QLineEdit(f["title"])
        self.e_body_font = QLineEdit(f["body"])
        self.e_h1_font = QLineEdit(f["h1"])
        self.e_h2_font = QLineEdit(f["h2"])
        self.s_title_size = self._dspin(s["title"], 0.5, 72)
        self.s_body_size = self._dspin(s["body"], 0.5, 72)
        self.s_title_line = self._dspin(l["title"], 1, 100)
        self.s_body_line = self._dspin(l["body"], 1, 100)
        self.s_top = self._dspin(pg["margin_top_mm"], 0, 100, 1)
        self.s_bottom = self._dspin(pg["margin_bottom_mm"], 0, 100, 1)
        self.s_left = self._dspin(pg["margin_left_mm"], 0, 100, 1)
        self.s_right = self._dspin(pg["margin_right_mm"], 0, 100, 1)

        form.addRow("标题字体", self.e_title_font)
        form.addRow("正文字体", self.e_body_font)
        form.addRow("一级标题字体", self.e_h1_font)
        form.addRow("二级标题字体", self.e_h2_font)
        form.addRow("标题字号(pt)", self.s_title_size)
        form.addRow("正文字号(pt)", self.s_body_size)
        form.addRow("标题行距(pt)", self.s_title_line)
        form.addRow("正文行距(pt)", self.s_body_line)
        form.addRow("上边距(mm)", self.s_top)
        form.addRow("下边距(mm)", self.s_bottom)
        form.addRow("左边距(mm)", self.s_left)
        form.addRow("右边距(mm)", self.s_right)

        btn_json = QPushButton("用文本编辑器打开 rules.json")
        btn_json.clicked.connect(self._open_json)
        form.addRow(btn_json)

        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.accepted.connect(self._on_ok)
        btns.rejected.connect(self.reject)
        form.addRow(btns)

    def _dspin(self, value, lo, hi, decimals=1):
        s = QDoubleSpinBox()
        s.setRange(lo, hi)
        s.setDecimals(decimals)
        s.setValue(float(value))
        return s

    def _open_json(self):
        import os
        p = rules_config.user_rules_path()
        if not p.exists():
            rules_config.save_rules(self.rules, p)
        os.startfile(str(p))

    def _on_ok(self):
        f = self.rules["fonts"]
        f["title"] = self.e_title_font.text().strip()
        f["body"] = self.e_body_font.text().strip()
        f["h1"] = self.e_h1_font.text().strip()
        f["h2"] = self.e_h2_font.text().strip()
        self.rules["sizes_pt"]["title"] = self.s_title_size.value()
        self.rules["sizes_pt"]["body"] = self.s_body_size.value()
        self.rules["line_spacing_pt"]["title"] = self.s_title_line.value()
        self.rules["line_spacing_pt"]["body"] = self.s_body_line.value()
        pg = self.rules["page"]
        pg["margin_top_mm"] = self.s_top.value()
        pg["margin_bottom_mm"] = self.s_bottom.value()
        pg["margin_left_mm"] = self.s_left.value()
        pg["margin_right_mm"] = self.s_right.value()
        self.accept()

    def get_rules(self):
        return self.rules


class RuleManagerDialog(QDialog):
    """规则管理：显示当前规则，可新增/删除角色，编辑后保存。"""

    def __init__(self, rules, parent=None):
        super().__init__(parent)
        self.setWindowTitle("规则管理")
        self.setMinimumSize(680, 480)
        self.rules = rules
        layout = QVBoxLayout(self)
        self.table = QTableWidget(0, 5)
        self.table.setHorizontalHeaderLabels(["角色", "字体", "字号(pt)", "行距(pt)", "对齐"])
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        layout.addWidget(self.table)
        btn_row = QHBoxLayout()
        add_btn = QPushButton("新增角色")
        add_btn.clicked.connect(self._add_role)
        del_btn = QPushButton("删除选中角色")
        del_btn.clicked.connect(self._del_role)
        btn_row.addWidget(add_btn)
        btn_row.addWidget(del_btn)
        btn_row.addStretch(1)
        layout.addLayout(btn_row)
        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.accepted.connect(self._on_ok)
        btns.rejected.connect(self.reject)
        layout.addWidget(btns)
        self._populate()

    def _populate(self):
        roles = [r for r in self.rules.get("fonts", {}).keys() if r != "ascii"]
        self.table.setRowCount(0)
        for role in roles:
            row = self.table.rowCount()
            self.table.insertRow(row)
            self.table.setItem(row, 0, QTableWidgetItem(role))
            self.table.setItem(row, 1, QTableWidgetItem(self.rules["fonts"].get(role, "")))
            self.table.setItem(row, 2, QTableWidgetItem(str(self.rules["sizes_pt"].get(role, ""))))
            self.table.setItem(row, 3, QTableWidgetItem(str(self.rules["line_spacing_pt"].get(role, ""))))
            self.table.setItem(row, 4, QTableWidgetItem(self.rules["alignment"].get(role, "")))

    def _add_role(self):
        name, ok = QInputDialog.getText(self, "新增角色", "角色名（英文键，如 custom）:")
        if ok and name.strip():
            role = name.strip()
            self.rules["fonts"].setdefault(role, "仿宋_GB2312")
            self.rules["sizes_pt"].setdefault(role, 16)
            self.rules["line_spacing_pt"].setdefault(role, 29)
            self.rules["alignment"].setdefault(role, "justify")
            self._populate()

    def _del_role(self):
        row = self.table.currentRow()
        if row < 0:
            QMessageBox.information(self, "提示", "请先选中要删除的角色")
            return
        role = self.table.item(row, 0).text()
        for d in ("fonts", "sizes_pt", "line_spacing_pt", "alignment"):
            self.rules[d].pop(role, None)
        self._populate()

    def _on_ok(self):
        for row in range(self.table.rowCount()):
            role = self.table.item(row, 0).text()
            font = self.table.item(row, 1).text().strip() if self.table.item(row, 1) else ""
            size = self.table.item(row, 2).text().strip() if self.table.item(row, 2) else ""
            line = self.table.item(row, 3).text().strip() if self.table.item(row, 3) else ""
            align = self.table.item(row, 4).text().strip() if self.table.item(row, 4) else ""
            if font:
                self.rules["fonts"][role] = font
            if size:
                try:
                    self.rules["sizes_pt"][role] = float(size)
                except ValueError:
                    pass
            if line:
                try:
                    self.rules["line_spacing_pt"][role] = float(line)
                except ValueError:
                    pass
            if align:
                self.rules["alignment"][role] = align
        self.accept()

    def get_rules(self):
        return self.rules


def _install_excepthook():
    import traceback

    def hook(etype, value, tb):
        try:
            if getattr(sys, "frozen", False):
                base = Path(sys.executable).resolve().parent
            else:
                base = Path(__file__).resolve().parent.parent
            with open(base / "公文格式自动排版助手_错误日志.txt", "a", encoding="utf-8") as f:
                f.write("".join(traceback.format_exception(etype, value, tb)))
                f.write("\n" + "=" * 40 + "\n")
        except Exception:
            pass
        sys.__excepthook__(etype, value, tb)

    sys.excepthook = hook


def main():
    _install_excepthook()
    app = QApplication(sys.argv)
    win = MainWindow()
    win.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
