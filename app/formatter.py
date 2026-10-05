"""格式应用引擎：按角色把规则套用到段落、页面与页码。"""
from __future__ import annotations

import unicodedata

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from . import rules_config
from .docx_io import (
    apply_page_setup,
    set_alignment,
    set_indent_chars,
    set_line_exact,
    set_run_font,
)

_INDENT_ROLES_FIRST_LINE = ("body", "h1", "h2", "h3", "h4")


def apply_formatting(document, roles: dict, rules: dict, date_text: str = None) -> None:
    """把规则套用到整个文档（页面、段落、空行、删除、日期替换、页码）。"""
    # 1) 页面设置（纵向节完整设置，横向节仅页边距）
    for section in document.sections:
        apply_page_setup(section, rules)

    # 2) 逐段套格式（段前段后=0、左右缩进=0、首行缩进按角色；date 替换日期文本）
    delete_paras = []
    for idx, p in enumerate(document.paragraphs):
        role = roles.get(idx)
        if role == "delete":
            delete_paras.append(p)
            continue
        if role is None:
            continue
        if not (p.text or "").strip():
            continue
        if role == "date" and date_text:
            _set_para_text(p, date_text)
        _apply_paragraph(p, role, rules)

    # 3) 空行管理（标题→主送 1 空、附件说明前 1 空、署名前 2 空）
    _manage_spacing(document, roles, rules)

    # 4) 删除标记段落
    for p in delete_paras:
        p._p.getparent().remove(p._p)

    # 5) 页码
    if rules.get("page_number", {}).get("enabled", True):
        _apply_page_numbers(document, rules)


def _set_para_text(p, text: str) -> None:
    """清空段落全部 run 并写入新文本。"""
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.add_run(text)


def _set_normal_style(document, rules):
    style = document.styles["Normal"]
    ascii_f = rules_config.ascii_font(rules)
    body_font = rules_config.role_font(rules, "body")
    body_size = rules_config.role_size_pt(rules, "body")
    style.font.name = ascii_f
    style.font.size = Pt(body_size)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), ascii_f)
    rFonts.set(qn("w:hAnsi"), ascii_f)
    rFonts.set(qn("w:eastAsia"), body_font)


def _apply_paragraph(p, role: str, rules: dict) -> None:
    size = rules_config.role_size_pt(rules, role)
    east = rules_config.role_font(rules, role)
    ascii_f = rules_config.ascii_font(rules)
    line = rules_config.role_line_spacing_pt(rules, role)
    align = rules_config.role_alignment(rules, role)

    set_alignment(p, align)
    set_line_exact(p, line)
    _apply_indent(p, role, rules, size)

    for run in p.runs:
        if _run_has_drawing(run):
            continue  # 图片/图形 run 不修改
        set_run_font(run, east, ascii_f, size)


def _run_has_drawing(run) -> bool:
    return bool(run._r.findall(".//" + qn("w:drawing"))) or \
        bool(run._r.findall(".//" + qn("w:pict")))


def _apply_indent(p, role: str, rules: dict, size: float) -> None:
    """首行缩进按角色；附件条目左缩进 5 字；署名右空 6 字、成文日期右空 5 字。"""
    ind = rules["indent"]
    first = None
    left = 0
    right = 0
    if role in _INDENT_ROLES_FIRST_LINE:
        first = ind.get("heading_first_line_chars", 2) if role.startswith("h") \
            else ind.get("body_first_line_chars", 2)
    elif role == "attachment_note":
        first = ind.get("attachment_first_line_chars", 2)
    elif role == "note":
        first = ind.get("note_first_line_chars", 2)
    elif role == "attachment_item":
        left = ind.get("attachment_item_left_chars", 5)
    elif role == "sign":
        right = ind.get("sign_right_chars", 6)
    elif role == "date":
        right = ind.get("date_right_chars", 5)
    set_indent_chars(p, size, first_line_chars=first, left_chars=left, right_chars=right)


def _manage_spacing(document, roles: dict, rules: dict) -> None:
    """空行管理：标题→主送 1 空；附件说明前 1 空；署名前 2 空；附件另面编排。"""
    paras = document.paragraphs
    by_role = {}
    for idx, p in enumerate(paras):
        role = roles.get(idx)
        if role and role not in by_role:
            by_role[role] = p

    body_line = rules_config.role_line_spacing_pt(rules, "body")
    if "title" in by_role and "addressee" in by_role:
        _ensure_empty_before(by_role["addressee"], 1, body_line)
    if "attachment_note" in by_role:
        _ensure_empty_before(by_role["attachment_note"], 1, body_line)
    if "sign" in by_role:
        _ensure_empty_before(by_role["sign"], 2, body_line)

    # 附件另面编排：附件前删除空行并加分页符，附件后空一行接标题
    if "attachment" in by_role:
        att = by_role["attachment"]
        _remove_empty_before(att)
        att._p.addprevious(_make_page_break_para())
        _ensure_empty_after(att, 1, body_line)


def _ensure_empty_before(paragraph, n: int, line_pt: float) -> None:
    """确保 paragraph 之前恰有 n 个连续空段（空段使用正文行距）。"""
    empties = []
    el = paragraph._p.getprevious()
    while el is not None and el.tag == qn("w:p") and _is_empty_p(el):
        empties.append(el)
        el = el.getprevious()
    current = len(empties)
    if current < n:
        for _ in range(n - current):
            paragraph._p.addprevious(_make_empty_para(line_pt))
    elif current > n:
        for e in empties[n:]:
            e.getparent().remove(e)
    # 归一化保留空段的行距
    kept = []
    el = paragraph._p.getprevious()
    while el is not None and el.tag == qn("w:p") and _is_empty_p(el) and len(kept) < n:
        kept.append(el)
        el = el.getprevious()
    for e in kept:
        _set_p_line_exact(e, line_pt)


def _ensure_empty_after(paragraph, n: int, line_pt: float) -> None:
    """确保 paragraph 之后恰有 n 个连续空段。"""
    empties = []
    el = paragraph._p.getnext()
    while el is not None and el.tag == qn("w:p") and _is_empty_p(el):
        empties.append(el)
        el = el.getnext()
    current = len(empties)
    if current < n:
        for _ in range(n - current):
            paragraph._p.addnext(_make_empty_para(line_pt))
    elif current > n:
        for e in empties[n:]:
            e.getparent().remove(e)
    kept = []
    el = paragraph._p.getnext()
    while el is not None and el.tag == qn("w:p") and _is_empty_p(el) and len(kept) < n:
        kept.append(el)
        el = el.getnext()
    for e in kept:
        _set_p_line_exact(e, line_pt)


def _remove_empty_before(paragraph) -> None:
    """删除 paragraph 之前的所有连续空段。"""
    el = paragraph._p.getprevious()
    while el is not None and el.tag == qn("w:p") and _is_empty_p(el):
        nxt = el.getprevious()
        el.getparent().remove(el)
        el = nxt


def _make_page_break_para():
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    return p


def _is_empty_p(el) -> bool:
    for t in el.findall(".//" + qn("w:t")):
        if (t.text or "").strip():
            return False
    if el.findall(".//" + qn("w:drawing")) or el.findall(".//" + qn("w:pict")):
        return False
    return True


def _make_empty_para(line_pt: float):
    p = OxmlElement("w:p")
    _set_p_line_exact(p, line_pt)
    return p


def _set_p_line_exact(el, line_pt: float) -> None:
    pPr = el.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        el.insert(0, pPr)
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(qn("w:line"), str(int(round(line_pt * 20))))
    spacing.set(qn("w:lineRule"), "exact")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")


def _apply_page_numbers(document, rules) -> None:
    pn = rules["page_number"]
    font = rules["fonts"].get("page_number", "宋体")
    size = rules["sizes_pt"].get("page_number", 14)
    ascii_f = font  # 页码西文也使用中文字体（宋体）
    pg = rules.get("page", {})
    text_width_twips = (pg.get("width_mm", 210) - pg.get("margin_left_mm", 28)
                        - pg.get("margin_right_mm", 26)) * 56.7

    _enable_even_odd(document)

    for section in document.sections:
        odd = section.footer
        odd.is_linked_to_previous = False
        even = section.even_page_footer
        even.is_linked_to_previous = False
        _build_page_number_footer(odd, pn, "odd", font, ascii_f, size, text_width_twips)
        _build_page_number_footer(even, pn, "even", font, ascii_f, size, text_width_twips)


def _enable_even_odd(document) -> None:
    try:
        document.settings.odd_and_even_pages_header_footer = True
    except Exception:
        settings = document.settings.element
        if settings.find(qn("w:evenAndOddHeaders")) is None:
            settings.append(OxmlElement("w:evenAndOddHeaders"))


def _build_page_number_footer(footer, pn, kind, font, ascii_f, size, text_width_twips) -> None:
    for p in list(footer.paragraphs):
        p._p.getparent().remove(p._p)

    para = footer.add_paragraph()
    align = pn.get("align", "justify")
    left_chars = pn.get("left_chars", 1)
    right_chars = pn.get("right_chars", 1)

    # pPr 子元素按 schema 顺序写入：tabs < spacing < ind < jc
    if kind == "odd":
        # 双面打印1（外侧）：奇数页页码靠右，用右制表位实现
        _add_right_tab_stop(para, text_width_twips, right_chars, size)
    _set_single_line_spacing(para, pn)
    set_indent_chars(para, size, left_chars=left_chars, right_chars=right_chars)
    set_alignment(para, align)

    if kind == "odd":
        _add_tab_run(para, font, ascii_f, size)

    _add_text_run(para, pn.get("dash_before", "\u2014"), font, ascii_f, size)
    _add_text_run(para, pn.get("space", " "), font, ascii_f, size)
    _add_page_field(para, font, ascii_f, size)
    _add_text_run(para, pn.get("space", " "), font, ascii_f, size)
    _add_text_run(para, pn.get("dash_after", "\u2014"), font, ascii_f, size)


def _set_single_line_spacing(para, pn) -> None:
    """单倍行距：w:line="240" w:lineRule="auto"（设置值为 1）。"""
    value = float(pn.get("line_spacing_value", 1.0))
    pPr = para._p.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(qn("w:line"), str(int(round(240 * value))))
    spacing.set(qn("w:lineRule"), "auto")


def _add_right_tab_stop(para, text_width_twips, right_chars, size) -> None:
    pos = text_width_twips - right_chars * size * 20
    pPr = para._p.get_or_add_pPr()
    tabs = pPr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        pPr.append(tabs)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(int(round(pos))))
    tabs.append(tab)


def _add_tab_run(para, font, ascii_f, size):
    run = para.add_run()
    run._r.append(OxmlElement("w:tab"))
    set_run_font(run, font, ascii_f, size)
    return run


def _add_text_run(para, text, font, ascii_f, size):
    run = para.add_run(text)
    set_run_font(run, font, ascii_f, size)
    return run


def _add_page_field(para, font, ascii_f, size) -> None:
    r_begin = para.add_run()
    _fld_char(r_begin, "begin")
    r_instr = para.add_run()
    _instr_text(r_instr, " PAGE ")
    r_sep = para.add_run()
    _fld_char(r_sep, "separate")
    r_value = para.add_run("1")
    r_end = para.add_run()
    _fld_char(r_end, "end")
    for run in (r_begin, r_instr, r_sep, r_value, r_end):
        set_run_font(run, font, ascii_f, size)


def _fld_char(run, type_: str) -> None:
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), type_)
    run._r.append(fld)


def _instr_text(run, text: str) -> None:
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = text
    run._r.append(instr)


# ---------- 对外暴露（供模板套用等复用） ----------
def format_paragraph(p, role: str, rules: dict) -> None:
    _apply_paragraph(p, role, rules)


def set_normal_style(document, rules) -> None:
    _set_normal_style(document, rules)


def apply_page_numbers(document, rules) -> None:
    _apply_page_numbers(document, rules)
