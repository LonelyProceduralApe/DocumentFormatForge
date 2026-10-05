"""综合测试：横向页边距 / 段前段后左右缩进0 / 空行管理 / 版头保留红线。"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / ".pylib"))

from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml.ns import qn

from app import classifier, formatter, rules_config, template_apply
from app.docx_io import read_paragraphs


def build_source():
    doc = Document()
    doc.add_paragraph("关于开展安全生产大检查的通知")      # 标题
    doc.add_paragraph("各乡镇人民政府、各街道办事处：")      # 主送
    doc.add_paragraph("为切实做好安全生产工作，现将有关事项通知如下。")  # 正文
    doc.add_paragraph("一、提高思想认识")                  # h1
    doc.add_paragraph("附件：1.检查清单")                  # 附件说明
    doc.add_paragraph("某某县人民政府")                    # 署名
    doc.add_paragraph("2024年1月5日")                      # 日期
    sec = doc.add_section(WD_SECTION.NEW_PAGE)             # 横向节
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    doc.add_paragraph("横向节内容")
    doc.save(str(ROOT / "_src.docx"))
    return doc


def section_report(doc):
    for i, s in enumerate(doc.sections):
        o = "横" if s.page_width > s.page_height else "纵"
        print(f"  节{i}: {o} 上{s.top_margin.mm:.0f} 下{s.bottom_margin.mm:.0f} 左{s.left_margin.mm:.0f} 右{s.right_margin.mm:.0f}")


def para_report(doc):
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if not t:
            continue
        ppr = p._p.pPr
        ind = sp = None
        if ppr is not None:
            ind = ppr.find(qn("w:ind"))
            sp = ppr.find(qn("w:spacing"))
        fl = ind.get(qn("w:firstLineChars")) if ind is not None else None
        l = ind.get(qn("w:leftChars")) if ind is not None else None
        r = ind.get(qn("w:rightChars")) if ind is not None else None
        before = sp.get(qn("w:before")) if sp is not None else None
        after = sp.get(qn("w:after")) if sp is not None else None
        print(f"  '{t[:14]}': firstLine={fl} left={l} right={r} before={before} after={after}")


def seq_report(doc):
    print("  序列:", " | ".join("(空)" if not (p.text or "").strip() else (p.text or "").strip()[:8] for p in doc.paragraphs))


def main():
    rules = rules_config.load_rules()

    doc = build_source()
    roles = classifier.classify(read_paragraphs(doc))
    formatter.apply_formatting(doc, roles, rules)
    doc.save(str(ROOT / "_fmt.docx"))
    v = Document(str(ROOT / "_fmt.docx"))
    print("=== 格式化：节 ===")
    section_report(v)
    print("=== 格式化：段落缩进/间距 ===")
    para_report(v)
    print("=== 格式化：段落序列 ===")
    seq_report(v)

    doc = build_source()
    roles = classifier.classify(read_paragraphs(doc))
    tpl = template_apply.list_templates()[0][1]
    out = template_apply.apply_template(tpl, doc, roles, rules,
                                        doc_num="白委〔2026〕12号", issuer_name="张三")
    out.save(str(ROOT / "_tpl.docx"))
    v = Document(str(ROOT / "_tpl.docx"))
    print("=== 模板：前若干段 ===")
    for p in v.paragraphs[:14]:
        t = (p.text or "").strip()
        pict = len(p._p.findall(".//" + qn("w:pict")))
        fonts = []
        for r in p.runs:
            rpr = r._r.rPr
            if rpr is not None:
                rf = rpr.find(qn("w:rFonts"))
                if rf is not None and rf.get(qn("w:eastAsia")):
                    fonts.append(ascii(rf.get(qn("w:eastAsia"))))
        print(f"  pict={pict} text={ascii(t[:20])} fonts={fonts}")


if __name__ == "__main__":
    main()
