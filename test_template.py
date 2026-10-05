"""测试版头模板套用。"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.oxml.ns import qn

from app import classifier, rules_config, template_apply
from app.docx_io import read_paragraphs

SAMPLE = ROOT / "sample_input.docx"


def build_sample():
    doc = Document()
    doc.add_paragraph("关于开展安全生产大检查的通知")
    doc.add_paragraph("各乡镇人民政府、各街道办事处：")
    doc.add_paragraph("为切实做好安全生产工作，现将有关事项通知如下。")
    doc.add_paragraph("一、提高思想认识")
    doc.add_paragraph("（一）加强组织领导")
    doc.add_paragraph("1.明确责任分工")
    doc.add_paragraph("（1）层层压实责任")
    doc.add_paragraph("附件：1.安全生产检查清单")
    doc.add_paragraph("某某县人民政府")
    doc.add_paragraph("2024年1月5日")
    doc.add_paragraph("（此件公开发布）")
    doc.save(str(SAMPLE))


def main():
    build_sample()
    rules = rules_config.load_rules()
    doc = Document(str(SAMPLE))
    roles = classifier.classify(read_paragraphs(doc))

    templates = template_apply.list_templates()
    print("发现模板:", [n for n, _ in templates])

    for name, path in templates:
        out_doc = template_apply.apply_template(path, doc, roles, rules)
        out = ROOT / f"sample_with_{name}.docx"
        out_doc.save(str(out))
        v = Document(str(out))
        print(f"\n=== {name} 输出结构 ===")
        for i, p in enumerate(v.paragraphs):
            has_img = bool(p._p.findall(".//" + qn("w:drawing")))
            t = (p.text or "").strip()
            if t or has_img:
                print(f"  [{i}] img={has_img} '{t[:28]}'")
            else:
                print(f"  [{i}] (空)")


if __name__ == "__main__":
    main()
