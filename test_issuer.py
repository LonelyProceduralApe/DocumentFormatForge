"""测试：发文字号/签发人占位提取与填写。"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from docx import Document

from app import classifier, rules_config, template_apply
from app.docx_io import read_paragraphs

SAMPLE = ROOT / "sample_input.docx"


def build_sample():
    doc = Document()
    doc.add_paragraph("关于开展安全生产大检查的通知")
    doc.add_paragraph("各乡镇人民政府、各街道办事处：")
    doc.add_paragraph("为切实做好安全生产工作，现将有关事项通知如下。")
    doc.add_paragraph("某某县人民政府")
    doc.add_paragraph("2024年1月5日")
    doc.save(str(SAMPLE))


def main():
    build_sample()
    rules = rules_config.load_rules()
    doc = Document(str(SAMPLE))
    roles = classifier.classify(read_paragraphs(doc))

    for name, path in template_apply.list_templates():
        fields = template_apply.extract_issuer_fields(path)
        print(f"[{name}] 提取占位: 发文字号={fields['doc_num']!r} 签发人={fields['name']!r}")

        out_doc = template_apply.apply_template(
            path, doc, roles, rules,
            doc_num=f"{name}〔2026〕12号", issuer_name="张三",
        )
        out = ROOT / f"sample_{name}_filled.docx"
        out_doc.save(str(out))
        v = Document(str(out))
        for p in v.paragraphs:
            if "签发人" in (p.text or "") or ("〔" in (p.text or "") and "号" in (p.text or "")):
                print(f"  [{name}] 填写后发文字号行: {p.text!r}")

    # 验证空签发人（下行文）场景
    out_doc = template_apply.apply_template(
        template_apply.list_templates()[0][1], doc, roles, rules,
        doc_num="白委〔2026〕8号", issuer_name="",
    )
    out = ROOT / "sample_noissuer.docx"
    out_doc.save(str(out))
    v = Document(str(out))
    for p in v.paragraphs:
        if "〔" in (p.text or ""):
            print(f"  [无签发人] 发文字号行: {p.text!r}")


if __name__ == "__main__":
    main()
