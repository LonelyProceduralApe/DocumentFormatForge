"""端到端冒烟测试：构造样例公文 → 识别 → 套格式 → 校验输出。"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.oxml.ns import qn

from app import classifier, formatter, rules_config
from app.docx_io import read_paragraphs

SAMPLE = ROOT / "sample_input.docx"
OUTPUT = ROOT / "sample_output.docx"


def build_sample():
    doc = Document()
    doc.add_paragraph("关于开展安全生产大检查的通知")          # 标题
    doc.add_paragraph("各乡镇人民政府、各街道办事处：")          # 主送机关
    doc.add_paragraph("为切实做好安全生产工作，现将有关事项通知如下。")  # 正文
    doc.add_paragraph("一、提高思想认识")                        # h1
    doc.add_paragraph("（一）加强组织领导")                       # h2
    doc.add_paragraph("1.明确责任分工")                           # h3
    doc.add_paragraph("（1）层层压实责任")                        # h4
    doc.add_paragraph("附件：1.安全生产检查清单")                 # 附件说明
    doc.add_paragraph("某某县人民政府")                           # 署名
    doc.add_paragraph("2024年1月5日")                             # 成文日期
    doc.add_paragraph("（此件公开发布）")                         # 附注
    doc.save(str(SAMPLE))
    print(f"样例已生成：{SAMPLE}")


def verify():
    doc = Document(str(OUTPUT))
    sec = doc.sections[0]
    print("\n=== 页面设置 ===")
    print(f"  页边距 上{sec.top_margin.mm:.1f} 下{sec.bottom_margin.mm:.1f} "
          f"左{sec.left_margin.mm:.1f} 右{sec.right_margin.mm:.1f} mm")
    print(f"  纸张 {sec.page_width.mm:.0f}×{sec.page_height.mm:.0f} mm")
    print("\n=== 段落格式（前几段）===")
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        if not t:
            continue
        r = p.runs[0] if p.runs else None
        east = None
        sz = None
        if r is not None and r._r.rPr is not None:
            rf = r._r.rPr.find(qn("w:rFonts"))
            if rf is not None:
                east = rf.get(qn("w:eastAsia"))
            szel = r._r.rPr.find(qn("w:sz"))
            if szel is not None:
                sz = int(szel.get(qn("w:val"))) / 2
        align = p.alignment
        ppr = p._p.pPr
        line = None
        if ppr is not None:
            sp = ppr.find(qn("w:spacing"))
            if sp is not None:
                line = sp.get(qn("w:line"))
        print(f"  [{i}] {t[:16]:<16} 字体={east} {sz}pt 对齐={align} 行距twips={line}")
        if i >= 9:
            break


def main():
    build_sample()
    rules = rules_config.load_rules()
    doc = Document(str(SAMPLE))
    paras = read_paragraphs(doc)
    roles = classifier.classify(paras)
    print("\n=== 识别结果 ===")
    for p in paras:
        if p["text"]:
            print(f"  [{p['idx']}] {roles.get(p['idx'], '跳过'):<16} {p['text'][:24]}")
    formatter.apply_formatting(doc, roles, rules)
    doc.save(str(OUTPUT))
    print(f"\n已保存：{OUTPUT}")
    verify()


if __name__ == "__main__":
    main()
