"""测试页码格式：两端对齐 + 前后各1字符 + 单倍行距 + 双面打印1（外侧）。"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.oxml.ns import qn

from app import classifier, formatter, rules_config
from app.docx_io import read_paragraphs


def main():
    doc = Document()
    doc.add_paragraph("关于开展安全生产大检查的通知")
    doc.add_paragraph("各乡镇人民政府、各街道办事处：")
    doc.add_paragraph("为切实做好安全生产工作，现将有关事项通知如下。")
    rules = rules_config.load_rules()
    roles = classifier.classify(read_paragraphs(doc))
    formatter.apply_formatting(doc, roles, rules)
    out = ROOT / "sample_pn.docx"
    doc.save(str(out))

    v = Document(str(out))
    sec = v.sections[0]
    for kind, footer in (("奇数页", sec.footer), ("偶数页", sec.even_page_footer)):
        print(f"=== {kind} 页脚 ===")
        for p in footer.paragraphs:
            ppr = p._p.pPr
            has_tab_run = any(r._r.find(qn("w:tab")) is not None for r in p.runs)
            print("  文本:", repr(p.text), "| 制表符run:", has_tab_run)
            print("  对齐:", p.alignment)
            if ppr is not None:
                ind = ppr.find(qn("w:ind"))
                if ind is not None:
                    print(f"  缩进 leftChars={ind.get(qn('w:leftChars'))} "
                          f"rightChars={ind.get(qn('w:rightChars'))} "
                          f"firstLineChars={ind.get(qn('w:firstLineChars'))}")
                sp = ppr.find(qn("w:spacing"))
                if sp is not None:
                    print(f"  行距 line={sp.get(qn('w:line'))} lineRule={sp.get(qn('w:lineRule'))}")
                tabs = ppr.find(qn("w:tabs"))
                if tabs is not None:
                    for t in tabs.findall(qn("w:tab")):
                        print(f"  制表位 val={t.get(qn('w:val'))} pos={t.get(qn('w:pos'))}")
                else:
                    print("  制表位: 无")


if __name__ == "__main__":
    main()
